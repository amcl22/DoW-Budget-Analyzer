"""Read the R-1 display workbook into raw string rows using the configured column map."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl

from .config import ExhibitColumns, normalize_header

HEADER_SEARCH_ROWS = 10


class ParseError(Exception):
    pass


@dataclass(frozen=True)
class RawRow:
    row_number: int               # 1-based row in the sheet
    fields: dict[str, str]        # canonical field -> cell text ('' when blank)
    amounts: dict[str, str]       # Excel header -> cell text ('' when blank)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_r1_workbook(path: Path, cols: ExhibitColumns) -> list[RawRow]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        if cols.sheet not in wb.sheetnames:
            raise ParseError(f"Sheet {cols.sheet!r} not found; sheets are {wb.sheetnames}")
        rows = list(wb[cols.sheet].iter_rows(values_only=True))
    finally:
        wb.close()

    # Locate the header row by content, not by a fixed index.
    wanted = {normalize_header(h) for h in cols.header_row_contains}
    header_idx = next(
        (
            i
            for i, row in enumerate(rows[:HEADER_SEARCH_ROWS])
            if wanted <= {normalize_header(c) for c in row if c is not None}
        ),
        None,
    )
    if header_idx is None:
        raise ParseError(
            f"No header row containing {sorted(cols.header_row_contains)} "
            f"in the first {HEADER_SEARCH_ROWS} rows of {cols.sheet!r}"
        )

    headers = [normalize_header(c) if c is not None else None for c in rows[header_idx]]
    position = {h: i for i, h in enumerate(headers) if h is not None}

    mapped = {normalize_header(h) for h in cols.fields.values()} | {
        normalize_header(a.header) for a in cols.amounts
    }
    missing = sorted(mapped - position.keys())
    unmapped = sorted(position.keys() - mapped)
    if missing or unmapped:
        raise ParseError(
            "Workbook headers do not match the column map. "
            f"Missing: {missing or 'none'}. Not in the column map: {unmapped or 'none'}. "
            f"Headers found: {[h for h in rows[header_idx] if h is not None]}"
        )

    field_pos = {f: position[normalize_header(h)] for f, h in cols.fields.items()}
    amount_pos = {a.header: position[normalize_header(a.header)] for a in cols.amounts}

    out = []
    for offset, row in enumerate(rows[header_idx + 1 :]):
        if all(c is None or _cell_text(c) == "" for c in row):
            continue
        cells = list(row) + [None] * (len(headers) - len(row))
        out.append(
            RawRow(
                row_number=header_idx + 2 + offset,
                fields={f: _cell_text(cells[i]) for f, i in field_pos.items()},
                amounts={h: _cell_text(cells[i]) for h, i in amount_pos.items()},
            )
        )
    return out
