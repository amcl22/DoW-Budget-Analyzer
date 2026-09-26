"""Rebuild the small offline test fixtures from the real releases.

    python tests/fixtures/make_fixtures.py [RAW_DIR]

RAW_DIR is the fetch store (default data/raw, filled by `budget ingest-all`). For each sample
this writes a handful of PDF pages and the Excel rows of the lines printed on them. All source
files are public DoW Comptroller releases.
"""

import sys
from pathlib import Path

import openpyxl
import pymupdf

from pipeline.config import load_accounts, load_column_map
from pipeline.exhibits import EXHIBITS
from pipeline.parse_xlsx import parse_workbook

HERE = Path(__file__).parent

SAMPLES = {
    # 6-7: department summaries (7 has 'FY 2026 Total*'); 10-11: Army BA 01-03 detail incl. a
    # classified line; 87: Golden Dome fund; 97: the DEFW agency section headed 'Defense-Wide'.
    "r1_sample": (2027, "R-1", [6, 7, 10, 11, 87, 97]),
    # older layout: 5 summary with adjacent 'Supplementals' headers ~9pt apart; 7 a row whose
    # label is missing from the text layer; 10 a stray '28' on the 'Applied Research' subtotal
    # and a first data row within the header band's depth.
    "r1_2024_sample": (2024, "R-1", [5, 7, 9, 10]),
    # rotated pages, columns split over page pairs: 6-9 Component Summary (A/B, grand total);
    # 14-15 2031A detail incl. AH-64 (gross, 'Less: Advance Procurement (PY)', net);
    # 160-161 1612N, whose line 1 has no printed line row.
    "p1_2027_sample": (2027, "P-1", [6, 7, 8, 9, 14, 15, 160, 161]),
    # single-page layout; 10: 'Advance Procurement (CY)' printed under the weapon line above;
    # 116: the same row separated from its line by a sub-activity heading.
    "p1_2024_sample": (2024, "P-1", [10, 116]),
}


def _source(raw: Path, fy: int, exhibit: str, prefix: str) -> Path:
    return next((raw / f"FY{fy}" / exhibit).glob(f"{prefix}*"))


def build(name: str, fy: int, exhibit: str, keep_pages: list[int], raw: Path) -> None:
    ex = EXHIBITS[exhibit]
    stem = exhibit.replace("-", "").lower()
    xlsx = _source(raw, fy, exhibit, f"{stem}_display")
    pdf = _source(raw, fy, exhibit, f"FY{fy}_{stem}")
    cols = load_column_map(exhibit, f"PB{fy}")
    records = ex.normalize(parse_workbook(xlsx, cols).rows, cols, load_accounts(), exhibit)
    links = ex.link(records, ex.parse_pdf(pdf, cols), cols)
    keep_rows = set()
    for r in records:
        if any(ref.page_number in keep_pages for ref in links.refs.get(r.key, [])):
            keep_rows.add(r.source_row_number)
            keep_rows.update(e.source_row_number for e in r.cost_elements)

    # Copy pages into a fresh document so resources used only by other pages (the 1.8 MB
    # cover seal) are not carried along.
    with pymupdf.open(pdf) as doc, pymupdf.open() as sample:
        for p in keep_pages:
            sample.insert_pdf(doc, from_page=p - 1, to_page=p - 1)
        sample.save(HERE / f"{name}.pdf", garbage=4, deflate=True)

    src = openpyxl.load_workbook(xlsx)[cols.sheet]
    out = openpyxl.Workbook()
    ws = out.active
    ws.title = cols.sheet
    for i, row in enumerate(src.iter_rows(values_only=True), start=1):
        if i <= 2 or i in keep_rows:
            ws.append(row)
    out.save(HERE / f"{name}.xlsx")
    print(f"{name}: {len(keep_pages)} pages, {len(keep_rows)} rows")


def build_r2(raw: Path) -> None:
    """First DTRA R-2 section (physical 23-29) and the first page of the next (30)."""
    book = next((raw / "FY2027" / "R-2").glob("RDTE_DTRA_PB_2027*"))
    with pymupdf.open(book) as doc, pymupdf.open() as sample:
        sample.insert_pdf(doc, from_page=22, to_page=29)
        sample.save(HERE / "r2_dtra_2027_sample.pdf", garbage=4, deflate=True)
    print("r2_dtra_2027_sample: 8 pages")


if __name__ == "__main__":
    raw = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("data/raw")
    for name, (fy, exhibit, pages) in SAMPLES.items():
        build(name, fy, exhibit, pages, raw)
    build_r2(raw)
