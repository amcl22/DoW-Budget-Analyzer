import openpyxl
import pytest

from pipeline.parse_r1 import ParseError, parse_r1_workbook
from tests.conftest import SAMPLE_XLSX


def test_finds_header_row_and_reads_every_line(raw_rows):
    assert len(raw_rows) == 50
    first = raw_rows[0]
    assert first.row_number == 3  # row 1 holds SUBTOTAL formulas, row 2 the headers
    assert first.fields["appropriation_account"] == "2040A"
    assert first.fields["program_element"] == "0601102A"
    assert first.amounts["FY 2025 Actuals"] == "290464"
    assert first.amounts["FY 2027 Mandatory Request"] == ""  # blank stays blank


def _modified_copy(tmp_path, edit):
    wb = openpyxl.load_workbook(SAMPLE_XLSX)
    edit(wb.active)
    path = tmp_path / "modified.xlsx"
    wb.save(path)
    return path


def test_unmapped_header_fails_the_run(tmp_path, cols):
    path = _modified_copy(tmp_path, lambda ws: ws.cell(row=2, column=ws.max_column + 1, value="FY 2028 Estimate"))
    with pytest.raises(ParseError, match="FY 2028 Estimate"):
        parse_r1_workbook(path, cols)


def test_missing_header_fails_the_run(tmp_path, cols):
    def rename(ws):
        for cell in ws[2]:
            if cell.value == "FY 2027 Total":
                cell.value = "FY 2027 Grand Total"
    path = _modified_copy(tmp_path, rename)
    with pytest.raises(ParseError, match="fy 2027 total"):
        parse_r1_workbook(path, cols)


def test_header_row_is_found_by_content_not_position(tmp_path, cols):
    path = _modified_copy(tmp_path, lambda ws: ws.insert_rows(1, amount=2))
    rows = parse_r1_workbook(path, cols)
    assert len(rows) == 50
    assert rows[0].row_number == 5
